import argparse
import sys
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
from sqlalchemy import select

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from app.db import Base, engine, SessionLocal
from app.models import Accommodation

def clean(v):
    return None if v is None or str(v).strip() == "" else str(v).strip()

def num(v, kind):
    try:
        return kind(v) if v not in (None, "") else None
    except (TypeError, ValueError):
        return None

def read_xlsx_standard(path):
    """Lecture autonome sans openpyxl via zipfile et XML standard."""
    with zipfile.ZipFile(path) as z:
        shared_strings = []
        if "xl/sharedStrings.xml" in z.namelist():
            ss_tree = ET.fromstring(z.read("xl/sharedStrings.xml"))
            for si in ss_tree.iter("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}si"):
                texts = [t.text for t in si.iter("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}t") if t.text]
                shared_strings.append("".join(texts))

        # Onglet 2 : Hebergements Occitanie
        sheet_xml = z.read("xl/worksheets/sheet2.xml")
        s_tree = ET.fromstring(sheet_xml)
        rows = []
        for row in s_tree.iter("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}row"):
            r_vals = []
            for c in row.iter("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}c"):
                t = c.attrib.get("t")
                is_node = c.find("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}is")
                if is_node is not None:
                    texts = [txt.text for txt in is_node.iter("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}t") if txt.text]
                    r_vals.append("".join(texts))
                else:
                    v = c.find("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}v")
                    val = v.text if v is not None else ""
                    if t == "s" and val.isdigit() and int(val) < len(shared_strings):
                        val = shared_strings[int(val)]
                    r_vals.append(val)
            rows.append(r_vals)
        return rows

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--file", required=True)
    args = parser.parse_args()
    path = Path(args.file)
    if not path.exists():
        raise SystemExit(f"Fichier introuvable : {path}")

    try:
        from openpyxl import load_workbook
        wb = load_workbook(path, data_only=True, read_only=True)
        if "Hebergements Occitanie" not in wb.sheetnames:
            raise SystemExit("Onglet Hebergements Occitanie absent.")
        ws = wb["Hebergements Occitanie"]
        rows = list(ws.iter_rows(values_only=True))
        wb.close()
    except ImportError:
        print("Note: openpyxl non installé, utilisation du parseur XML standard ultra-rapide.")
        rows = read_xlsx_standard(path)

    if not rows:
        raise SystemExit("Le fichier Excel est vide.")

    headers = [clean(v) for v in rows[0]]
    ix = {h: i for i, h in enumerate(headers) if h}

    Base.metadata.create_all(bind=engine)
    db = SessionLocal()
    counts = {"read": 0, "created": 0, "updated": 0, "rejected": 0}

    try:
        for rownum, row in enumerate(rows[1:], start=2):
            if not any(v is not None and str(v).strip() != "" for v in row):
                continue
            counts["read"] += 1
            def get(k):
                return row[ix[k]] if k in ix and ix[k] < len(row) else None

            name, commune = clean(get("Structure")), clean(get("Commune"))
            if not name or not commune:
                counts["rejected"] += 1
                continue

            source_id = f"{path.name}:Hebergements Occitanie:{rownum}"
            obj = db.scalar(select(Accommodation).where(Accommodation.source_record_id == source_id))
            data = {
                "source_record_id": source_id,
                "name": name,
                "commune": commune,
                "department": clean(get("Departement")),
                "postal_code": clean(get("CodePostal")),
                "latitude": num(get("Latitude"), float),
                "longitude": num(get("Longitude"), float),
                "themes_raw": clean(get("Thematiques")),
                "student_beds": num(get("Lits_eleves"), int),
                "supervisor_beds": num(get("Lits_encadrants"), int),
                "reference_tariff": clean(get("Tarif_reference_ministere")),
                "nature_label": clean(get("Label_ma_classe_nature")),
                "accessible_handicap": clean(get("Accessible_handicap")),
                "opening_period": clean(get("Periode_ouverture")),
                "classrooms": clean(get("Salles_de_classe")),
                "coach_charter": clean(get("Transport_autocar_charte")),
                "nearby_sites_raw": clean(get("Sites_a_proximite")),
                "website_url": clean(get("Site_internet")),
                "permalink": clean(get("Permalien")),
                "total_capacity": num(get("Capacite_totale"), int),
                "nearest_station_raw": clean(get("Gare_principale_plus_proche")),
                "straight_line_distance_km": num(get("Distance_vol_oiseau_km"), float),
                "source_name": path.name
            }
            if obj:
                for k, v in data.items():
                    setattr(obj, k, v)
                counts["updated"] += 1
            else:
                db.add(Accommodation(**data))
                counts["created"] += 1

        db.commit()
        print("Import terminé avec succès :", counts)
    finally:
        db.close()

if __name__ == "__main__":
    main()
